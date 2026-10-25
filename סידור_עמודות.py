"""
סידור עמודות אקסל - ניטור יתושים.

סדר העמודות הנכון (מקובץ Report Run) קבוע בתוך התוכנה.
בוחרים או גוררים קובץ אחד או יותר, והתוכנה שומרת לצד כל קובץ
עותק מסודר בשם "<שם הקובץ> - מסודר.xlsx". הקובץ המקורי לא משתנה.
"""

import os
import sys
import math
import traceback
from copy import copy
import tkinter as tk
from tkinter import filedialog, messagebox

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
except ImportError:
    tk.Tk().withdraw()
    messagebox.showerror("חסרה ספרייה", "יש להתקין את openpyxl:\n\npip install openpyxl")
    raise SystemExit(1)

try:
    from tkinterdnd2 import TkinterDnD, DND_FILES
    HAS_DND = True
except ImportError:
    HAS_DND = False


# ---------------------------------------------------------------- סדר העמודות (מ-Report Run)

COLUMN_ORDER = [
    ("WNV No.", 9.2),
    ("מיקום מבחנה בקופסה", 19.1),
    ("מזהה מבחנה", 28.2),
    ("תאריך איסוף", 11.1),
    ("שם נקודת ניטור", 58.0),
    ("שם אתר", 13.4),
    ("רשות שיפוט", 11.0),
    ("מחוז", 11.5),
    ("בדיקות נגיף - מין", 16.4),
    ("בדיקות נגיף - כמות נקבות", 21.9),
    ("מזהה ביקור", 11.2),
    ("Point X ITM", 11.1),
    ("Point Y ITM", 13.0),
    ("שם דוגם", 11.2),
    ("מזהה נקודת ניטור", 15.6),
    ("Sent/ Received at Sheba", 22.9),
    ("Date mosquitoes homogenized", 29.1),
    ("Date RNA extracted", 18.6),
    ("Date of TaqMan", 15.1),
    ("Result RT-PCR (Lin1+2)", 21.0),
]

HEADER_SCAN_ROWS = 15

# בהרצה כקובץ exe: התיקייה של ה-exe, והאייקון נטען מתוך ה-exe עצמו
FROZEN = getattr(sys, "frozen", False)
APP_DIR = os.path.dirname(sys.executable if FROZEN else os.path.abspath(__file__))
RESOURCE_DIR = getattr(sys, "_MEIPASS", APP_DIR)


# ---------------------------------------------------------------- לוגיקה

def norm(value):
    """מנרמל כותרת להשוואה: מוריד רווחים כפולים ומתעלם מאותיות גדולות/קטנות."""
    if value is None:
        return ""
    return " ".join(str(value).split()).casefold()


def find_header_row(ws, template_set):
    """שורת הכותרות = השורה עם הכי הרבה כותרות שמופיעות בסדר הנכון."""
    best_row, best_count = None, 0
    for r in range(1, min(ws.max_row, HEADER_SCAN_ROWS) + 1):
        count = sum(1 for c in range(1, ws.max_column + 1)
                    if norm(ws.cell(row=r, column=c).value) in template_set)
        if count > best_count:
            best_row, best_count = r, count
    return best_row, best_count


def reorder_sheet(ws):
    """מסדר את העמודות בגיליון לפי COLUMN_ORDER. מחזיר דוח, או None אם הגיליון לא רלוונטי."""
    template_set = {norm(h) for h, _ in COLUMN_ORDER}
    hr, matches = find_header_row(ws, template_set)
    if hr is None or matches < 2:
        return None

    max_col, max_row = ws.max_column, ws.max_row
    src_headers = {c: norm(ws.cell(row=hr, column=c).value) for c in range(1, max_col + 1)}

    # בניית הסדר החדש: (עמודת מקור או None, כותרת לעמודה חסרה, רוחב מהתבנית)
    used = set()
    plan, missing = [], []
    for header, width in COLUMN_ORDER:
        key = norm(header)
        src = next((c for c in range(1, max_col + 1)
                    if c not in used and src_headers[c] == key), None)
        if src is not None:
            used.add(src)
            plan.append((src, None, width))
        else:
            plan.append((None, header, width))
            missing.append(header)

    # עמודות שקיימות בקובץ אבל לא בסדר הנכון - נשמרות בסוף
    extra = []
    for c in range(1, max_col + 1):
        if c in used:
            continue
        if any(ws.cell(row=r, column=c).value is not None for r in range(hr, max_row + 1)):
            plan.append((c, None, None))
            extra.append(str(ws.cell(row=hr, column=c).value or get_column_letter(c)))

    moved = sum(1 for i, (src, _, _) in enumerate(plan, start=1) if src is not None and src != i)

    # צילום מצב של כל התאים מהכותרת ומטה
    for rng in list(ws.merged_cells.ranges):
        if rng.max_row >= hr:
            ws.unmerge_cells(str(rng))

    snapshot = {}
    for c in range(1, max_col + 1):
        snapshot[c] = [(cell.value, copy(cell._style), copy(cell.hyperlink) if cell.hyperlink else None)
                       for cell in (ws.cell(row=r, column=c) for r in range(hr, max_row + 1))]
    widths = {c: ws.column_dimensions[get_column_letter(c)].width for c in range(1, max_col + 1)}
    header_style = copy(ws.cell(row=hr, column=min(used) if used else 1)._style)

    # ניקוי האזור
    for c in range(1, max(max_col, len(plan)) + 1):
        for r in range(hr, max_row + 1):
            cell = ws.cell(row=r, column=c)
            cell.value = None
            cell.hyperlink = None
            cell.style = "Normal"

    # כתיבה בסדר החדש
    for new_c, (src, new_header, width) in enumerate(plan, start=1):
        letter = get_column_letter(new_c)
        if src is not None:
            for i, (value, style, link) in enumerate(snapshot[src]):
                cell = ws.cell(row=hr + i, column=new_c)
                cell.value = value
                cell._style = copy(style)
                if link:
                    cell.hyperlink = link
            if widths[src]:
                ws.column_dimensions[letter].width = widths[src]
        else:
            cell = ws.cell(row=hr, column=new_c)
            cell.value = new_header
            cell._style = copy(header_style)
            ws.column_dimensions[letter].width = width

    if ws.auto_filter.ref:
        ws.auto_filter.ref = f"A{hr}:{get_column_letter(len(plan))}{max_row}"

    return {"moved": moved, "missing": missing, "extra": extra}


def output_path_for(path):
    base, ext = os.path.splitext(path)
    return f"{base} - מסודר{ext}"


def process_file(path):
    """מסדר קובץ אחד ושומר עותק. מחזיר (סטטוס, הודעה)."""
    wb = openpyxl.load_workbook(path)
    reports = [r for r in (reorder_sheet(ws) for ws in wb.worksheets) if r]
    if not reports:
        return "err", "לא נמצאו כותרות מתאימות בקובץ"
    wb.save(output_path_for(path))

    moved = sum(r["moved"] for r in reports)
    missing = [m for r in reports for m in r["missing"]]
    extra = [e for r in reports for e in r["extra"]]
    notes = []
    if missing:
        notes.append("נוספו ריקות: " + ", ".join(missing))
    if extra:
        notes.append("הועברו לסוף: " + ", ".join(extra))
    if moved == 0 and not missing:
        msg = "העמודות כבר היו מסודרות"
    else:
        msg = f"הוזזו {moved} עמודות למקומן"
    if notes:
        return "warn", "\n".join([msg] + notes)
    return "ok", msg


# ---------------------------------------------------------------- עיצוב

BG = "#f6f7fb"
SURFACE = "#ffffff"
SURFACE_2 = "#f0f2f7"
BORDER = "#e3e6ee"
TEXT = "#1b1e28"
MUTED = "#6b7385"
FAINT = "#a3a9b8"
ACCENT = "#6c4cf5"
ACCENT_HOVER = "#5a3be3"
ACCENT_SOFT = "#efebff"
ACCENT_2 = "#4f8cff"
FONT = "Segoe UI"

SCALE = 1.0  # מתעדכן לפי הגדלת התצוגה של Windows


def px(v):
    return int(round(v * SCALE))


BADGES = {
    "wait": ("ממתין", "#eef0f5", "#5f6779"),
    "run":  ("מעבד…", "#e3f1ff", "#1f74c9"),
    "ok":   ("סודר ✓", "#e3f8ec", "#16834a"),
    "warn": ("סודר ✓", "#fff4dc", "#b26a00"),
    "err":  ("שגיאה", "#fde8e8", "#c62828"),
}


def round_rect(canvas, x1, y1, x2, y2, r, **kw):
    pts = [x1 + r, y1, x2 - r, y1, x2, y1, x2, y1 + r, x2, y2 - r, x2, y2,
           x2 - r, y2, x1 + r, y2, x1, y2, x1, y2 - r, x1, y1 + r, x1, y1]
    return canvas.create_polygon(pts, smooth=True, **kw)


def hex_mix(a, b, t):
    a = [int(a[i:i + 2], 16) for i in (1, 3, 5)]
    b = [int(b[i:i + 2], 16) for i in (1, 3, 5)]
    return "#" + "".join(f"{round(x + (y - x) * t):02x}" for x, y in zip(a, b))


class RoundButton(tk.Canvas):
    """כפתור מעוגל. style: primary / ghost."""

    def __init__(self, parent, text, command, width=150, height=36, style="primary",
                 parent_bg=BG, font_size=10):
        width, height = px(width), px(height)
        super().__init__(parent, width=width, height=height, bg=parent_bg,
                         highlightthickness=0, cursor="hand2")
        if style == "primary":
            self.colors = {"bg": ACCENT, "hover": ACCENT_HOVER, "fg": "white", "outline": ""}
        else:
            self.colors = {"bg": parent_bg, "hover": SURFACE_2, "fg": TEXT, "outline": BORDER}
        self.command = command
        self.enabled = True
        self.shape = round_rect(self, 1, 1, width - 1, height - 1, px(12), fill=self.colors["bg"],
                                outline=self.colors["outline"], width=1)
        self.label = self.create_text(width / 2, height / 2, text=text, fill=self.colors["fg"],
                                      font=(FONT, font_size, "bold"))
        self.bind("<Enter>", lambda e: self.enabled and self.itemconfig(self.shape, fill=self.colors["hover"]))
        self.bind("<Leave>", lambda e: self.enabled and self.itemconfig(self.shape, fill=self.colors["bg"]))
        self.bind("<Button-1>", lambda e: self.enabled and self.command())

    def set_enabled(self, enabled):
        self.enabled = enabled
        self.itemconfig(self.shape, fill=self.colors["bg"] if enabled else SURFACE_2)
        self.itemconfig(self.label, fill=self.colors["fg"] if enabled else FAINT)
        self.config(cursor="hand2" if enabled else "arrow")


def _ellipse(cx, cy, rx, ry, angle=0.0, n=40):
    a = math.radians(angle)
    pts = []
    for i in range(n):
        t = 2 * math.pi * i / n
        x, y = rx * math.cos(t), ry * math.sin(t)
        pts.append((cx + x * math.cos(a) - y * math.sin(a), cy + x * math.sin(a) + y * math.cos(a)))
    return pts


def mosquito_shapes():
    """יתוש במבט מלמעלה, ביחידות שבהן הגובה הכולל הוא בערך 1 והמרכז ב-(0,0)."""
    shapes = []
    # כנפיים (מתחת לגוף)
    for s in (-1, 1):
        shapes.append(("poly", _ellipse(s * 0.13, 0.03, 0.065, 0.21, angle=-s * 28), "wing"))
    # רגליים
    legs = [
        [(0.04, -0.19), (0.17, -0.27), (0.27, -0.42)],
        [(0.05, -0.15), (0.24, -0.13), (0.40, -0.02)],
        [(0.04, -0.11), (0.17, 0.04), (0.29, 0.40)],
    ]
    for s in (-1, 1):
        for leg in legs:
            shapes.append(("line", [(s * x, y) for x, y in leg], 0.022, "body"))
        shapes.append(("line", [(s * 0.015, -0.31), (s * 0.08, -0.40)], 0.018, "body"))  # מחוש
    # חדק
    shapes.append(("line", [(0, -0.30), (0, -0.50)], 0.02, "body"))
    # גוף
    shapes.append(("poly", _ellipse(0, 0.12, 0.045, 0.22), "body"))      # בטן
    shapes.append(("poly", _ellipse(0, -0.16, 0.072, 0.085), "body"))    # חזה
    shapes.append(("poly", _ellipse(0, -0.275, 0.042, 0.042), "body"))   # ראש
    return shapes


def draw_mosquito(c, cx, cy, size, body, wing):
    """מצייר יתוש על Canvas. size = הגובה בפיקסלים."""
    for kind, pts, *rest in mosquito_shapes():
        flat = [v for x, y in pts for v in (cx + x * size, cy + y * size)]
        if kind == "poly":
            c.create_polygon(flat, fill=wing if rest[-1] == "wing" else body, outline="")
        else:
            c.create_line(flat, fill=body, width=max(1, rest[0] * size), capstyle="round",
                          joinstyle="round")


def logo(parent, size=34, bg=BG):
    """לוגו: ריבוע מעוגל בגרדיאנט ועליו יתוש לבן."""
    s = px(size)
    c = tk.Canvas(parent, width=s, height=s, bg=bg, highlightthickness=0)
    for i in range(s):
        c.create_line(i, 0, i, s, fill=hex_mix(ACCENT_2, ACCENT, i / s))
    r = px(9)
    for (x0, y0, start) in ((0, 0, 90), (s - 2 * r, 0, 0), (0, s - 2 * r, 180), (s - 2 * r, s - 2 * r, 270)):
        # מסכה לפינות: ריבוע בצבע הרקע ועליו רבע עיגול בצבע הגרדיאנט
        cx, cy = x0 + r, y0 + r
        sx = -1 if start in (90, 180) else 1
        sy = -1 if start in (0, 90) else 1
        c.create_rectangle(cx, cy, cx + sx * r, cy + sy * r, fill=bg, outline="")
        c.create_arc(x0, y0, x0 + 2 * r, y0 + 2 * r, start=start, extent=90, style="pieslice",
                     fill=hex_mix(ACCENT_2, ACCENT, cx / s), outline="")
    draw_mosquito(c, s / 2, s * 0.53, s * 0.8, "white", hex_mix("#ffffff", ACCENT, 0.35))
    return c


def create_shortcut(lnk_path, target, args, workdir, icon, description):
    """יוצר קובץ lnk דרך IShellLinkW (תומך בנתיבים בעברית, בניגוד ל-WScript.Shell)."""
    import ctypes
    import uuid
    from ctypes import c_void_p, c_wchar_p, c_int, POINTER, byref, HRESULT, WINFUNCTYPE

    def guid(s):
        return (ctypes.c_byte * 16).from_buffer_copy(uuid.UUID(s).bytes_le)

    def call(obj, index, argtypes, *args):
        vtable = ctypes.cast(obj, POINTER(POINTER(c_void_p))).contents
        return WINFUNCTYPE(HRESULT, c_void_p, *argtypes)(vtable[index])(obj, *args)

    ole32 = ctypes.windll.ole32
    ole32.CoInitialize(None)
    link, persist = c_void_p(), c_void_p()
    ole32.CoCreateInstance(byref(guid("00021401-0000-0000-C000-000000000046")), None, 1,
                           byref(guid("000214F9-0000-0000-C000-000000000046")), byref(link))
    try:
        call(link, 20, [c_wchar_p], target)               # SetPath
        call(link, 11, [c_wchar_p], args)                 # SetArguments
        call(link, 9, [c_wchar_p], workdir)               # SetWorkingDirectory
        call(link, 17, [c_wchar_p, c_int], icon, 0)       # SetIconLocation
        call(link, 7, [c_wchar_p], description)           # SetDescription
        call(link, 0, [c_void_p, POINTER(c_void_p)],      # QueryInterface(IPersistFile)
             byref(guid("0000010B-0000-0000-C000-000000000046")), byref(persist))
        call(persist, 6, [c_wchar_p, c_int], lnk_path, 1)  # Save
    finally:
        for obj in (persist, link):
            if obj:
                vtable = ctypes.cast(obj, POINTER(POINTER(c_void_p))).contents
                WINFUNCTYPE(ctypes.c_ulong, c_void_p)(vtable[2])(obj)


# ---------------------------------------------------------------- ממשק

BaseTk = TkinterDnD.Tk if HAS_DND else tk.Tk


class App(BaseTk):
    def __init__(self):
        super().__init__()
        global SCALE
        SCALE = self.winfo_fpixels("1i") / 96
        self.title("סידור עמודות")
        w, h = px(640), px(470)
        self.geometry(f"{w}x{h}+{(self.winfo_screenwidth() - w) // 2}+{(self.winfo_screenheight() - h) // 3}")
        self.minsize(px(560), px(420))
        self.configure(bg=BG)
        self._set_icon()
        self.after(800, self._ensure_shortcut)

        self.files = []      # נתיבים לפי סדר הוספה
        self.rows = {}       # נתיב -> רכיבי השורה ברשימה
        self.popup = None

        wrap = tk.Frame(self, bg=BG)
        wrap.pack(fill="both", expand=True, padx=px(18), pady=px(14))
        self._build_topbar(wrap)
        self._build_dropzone(wrap)
        self._build_actions(wrap)
        self._build_file_list(wrap)
        self._refresh()
        self.bind("<Configure>", lambda e: e.widget is self and self._close_popup(), add="+")

    def _ensure_shortcut(self):
        """יוצר/מעדכן קיצור דרך עם סמל היתוש ליד התוכנה (מתעדכן גם אחרי העברה למחשב אחר)."""
        if FROZEN:
            return  # ל-exe יש אייקון מובנה, אין צורך בקיצור דרך
        folder = APP_DIR
        ico = os.path.join(folder, "mosquito.ico")
        pythonw = os.path.join(os.path.dirname(sys.executable), "pythonw.exe")
        if os.path.exists(ico) and os.path.exists(pythonw):
            try:
                create_shortcut(os.path.join(folder, "סידור עמודות.lnk"), pythonw,
                                f'"{os.path.abspath(__file__)}"', folder, ico, "סידור עמודות אקסל")
            except Exception:
                traceback.print_exc()

    def _set_icon(self):
        try:
            from ctypes import windll
            windll.shell32.SetCurrentProcessExplicitAppUserModelID("mosquito.column.sorter")
        except Exception:
            pass
        ico = os.path.join(RESOURCE_DIR, "mosquito.ico")
        if os.path.exists(ico):
            try:
                self.iconbitmap(default=ico)
            except tk.TclError:
                pass

    # --- שורה עליונה
    def _build_topbar(self, parent):
        bar = tk.Frame(parent, bg=BG)
        bar.pack(fill="x", pady=(0, px(12)))
        logo(bar).pack(side="right")
        titles = tk.Frame(bar, bg=BG)
        titles.pack(side="right", padx=px(10))
        tk.Label(titles, text="סידור עמודות", bg=BG, fg=TEXT,
                 font=(FONT, 13, "bold")).pack(anchor="e")
        tk.Label(titles, text="ניטור יתושים  ·  לפי סדר דוח הריצה", bg=BG, fg=MUTED,
                 font=(FONT, 8)).pack(anchor="e")
        self.order_btn = RoundButton(bar, "סדר העמודות", self.toggle_order, width=112, height=30,
                                     style="ghost", font_size=9)
        self.order_btn.pack(side="left")

    # --- חלון קופץ עם סדר העמודות
    def _close_popup(self):
        if self.popup is not None and self.popup.winfo_exists():
            self.popup.destroy()
        self.popup = None

    def toggle_order(self):
        if self.popup is not None:
            self._close_popup()
            return
        pop = tk.Toplevel(self, bg=BORDER)
        pop.overrideredirect(True)
        inner = tk.Frame(pop, bg=SURFACE)
        inner.pack(fill="both", expand=True, padx=1, pady=1)
        tk.Label(inner, text="סדר העמודות הקבוע", bg=SURFACE, fg=TEXT,
                 font=(FONT, 10, "bold")).pack(anchor="e", padx=px(14), pady=(px(10), px(6)))
        grid = tk.Frame(inner, bg=SURFACE)
        grid.pack(padx=px(10), pady=(0, px(10)))
        half = (len(COLUMN_ORDER) + 1) // 2
        for i, (name, _) in enumerate(COLUMN_ORDER):
            col, row = (1, i) if i < half else (0, i - half)   # העמודה הימנית קודם
            cell = tk.Frame(grid, bg=SURFACE)
            cell.grid(row=row, column=col, sticky="e", padx=px(6), pady=px(1))
            tk.Label(cell, text=str(i + 1), bg=ACCENT_SOFT, fg=ACCENT_HOVER, width=3,
                     font=(FONT, 7, "bold")).pack(side="right")
            tk.Label(cell, text=name, bg=SURFACE, fg=TEXT, font=(FONT, 8)).pack(side="right", padx=(0, px(6)))
        pop.update_idletasks()
        x = self.order_btn.winfo_rootx()
        y = self.order_btn.winfo_rooty() + self.order_btn.winfo_height() + px(6)
        pop.geometry(f"+{x}+{y}")
        pop.bind("<Button-1>", lambda e: self._close_popup())
        self.popup = pop

    # --- אזור גרירה
    def _build_dropzone(self, parent):
        self.drop = tk.Canvas(parent, height=px(84), bg=BG, highlightthickness=0, cursor="hand2")
        self.drop.pack(fill="x")
        self.drop_hover = False
        self.drop.bind("<Configure>", lambda e: self._draw_drop())
        self.drop.bind("<Enter>", lambda e: self._set_drop_hover(True))
        self.drop.bind("<Leave>", lambda e: self._set_drop_hover(False))
        self.drop.bind("<Button-1>", lambda e: self.pick_files())
        if HAS_DND:
            for w in (self.drop, self):
                w.drop_target_register(DND_FILES)
                w.dnd_bind("<<Drop>>", self._on_drop)
            self.drop.dnd_bind("<<DropEnter>>", lambda e: self._set_drop_hover(True))
            self.drop.dnd_bind("<<DropLeave>>", lambda e: self._set_drop_hover(False))

    def _set_drop_hover(self, value):
        self.drop_hover = value
        self._draw_drop()
        return "copy"

    def _draw_drop(self):
        c = self.drop
        c.delete("all")
        w, h = c.winfo_width(), c.winfo_height()
        hover = self.drop_hover
        round_rect(c, 1, 1, w - 1, h - 1, px(16), fill=ACCENT_SOFT if hover else SURFACE,
                   outline=ACCENT if hover else BORDER, dash=(5, 4), width=max(1, px(1)))
        # אייקון העלאה
        cx, cy, r = w - px(46), h / 2, px(20)
        c.create_oval(cx - r, cy - r, cx + r, cy + r, fill=ACCENT if hover else ACCENT_SOFT, outline="")
        col = "white" if hover else ACCENT_HOVER
        c.create_polygon(cx, cy - px(10), cx - px(7.5), cy - px(2), cx - px(3), cy - px(2),
                         cx - px(3), cy + px(6), cx + px(3), cy + px(6), cx + px(3), cy - px(2),
                         cx + px(7.5), cy - px(2), fill=col, outline="")
        c.create_line(cx - px(8), cy + px(10), cx + px(8), cy + px(10), fill=col,
                      width=px(2), capstyle="round")
        main = "גרור לכאן קבצי אקסל" if HAS_DND else "בחר קבצי אקסל"
        c.create_text(w - px(80), cy - px(9), text=main, anchor="e", fill=TEXT, font=(FONT, 11, "bold"))
        sub = "או לחץ לבחירה  ·  אפשר כמה קבצים יחד" if HAS_DND else "לחץ כאן  ·  אפשר כמה קבצים יחד"
        c.create_text(w - px(80), cy + px(11), text=sub, anchor="e", fill=MUTED, font=(FONT, 8))

    # --- רשימת קבצים
    def _build_file_list(self, parent):
        head = tk.Frame(parent, bg=BG)
        head.pack(fill="x", pady=(px(14), px(6)))
        tk.Label(head, text="קבצים", bg=BG, fg=MUTED, font=(FONT, 9, "bold")).pack(side="right")
        self.count_lbl = tk.Label(head, bg=SURFACE_2, fg=TEXT, font=(FONT, 7, "bold"), padx=px(6))
        self.count_lbl.pack(side="right", padx=px(6))
        self.clear_lbl = tk.Label(head, text="נקה הכל", bg=BG, fg=MUTED, font=(FONT, 8), cursor="hand2")
        self.clear_lbl.pack(side="left")
        self.clear_lbl.bind("<Button-1>", lambda e: self.clear())
        self.clear_lbl.bind("<Enter>", lambda e: self.clear_lbl.config(fg=TEXT))
        self.clear_lbl.bind("<Leave>", lambda e: self.clear_lbl.config(fg=MUTED))

        self.list_canvas = tk.Canvas(parent, bg=BG, highlightthickness=0)
        self.list_canvas.pack(fill="both", expand=True)
        self.list_inner = tk.Frame(self.list_canvas, bg=BG)
        win = self.list_canvas.create_window(0, 0, window=self.list_inner, anchor="nw")
        self.list_inner.bind("<Configure>", lambda e: self.list_canvas.configure(
            scrollregion=self.list_canvas.bbox("all")))
        self.list_canvas.bind("<Configure>", lambda e: self.list_canvas.itemconfig(win, width=e.width))
        self.list_canvas.bind_all("<MouseWheel>", self._on_wheel)

        self.empty_lbl = tk.Frame(self.list_inner, bg=BG)
        art = tk.Canvas(self.empty_lbl, width=px(56), height=px(56), bg=BG, highlightthickness=0)
        draw_mosquito(art, px(28), px(30), px(52), "#cdd1dc", "#e6e8f0")
        art.pack()
        tk.Label(self.empty_lbl, text="עדיין לא נבחרו קבצים", bg=BG, fg=FAINT,
                 font=(FONT, 9)).pack(pady=(px(4), 0))

    def _on_wheel(self, e):
        if self.list_inner.winfo_height() > self.list_canvas.winfo_height():
            self.list_canvas.yview_scroll(int(-e.delta / 120), "units")

    def _add_row(self, path):
        row = tk.Frame(self.list_inner, bg=SURFACE, highlightbackground=BORDER, highlightthickness=1)
        row.pack(fill="x", pady=(0, px(6)))

        icon = tk.Canvas(row, width=px(26), height=px(26), bg=SURFACE, highlightthickness=0)
        round_rect(icon, 0, 0, px(26), px(26), px(7), fill="#1f7a4d", outline="")
        icon.create_text(px(13), px(13), text="X", fill="white", font=(FONT, 9, "bold"))
        icon.pack(side="right", padx=(0, px(10)), pady=px(9), anchor="n")

        texts = tk.Frame(row, bg=SURFACE)
        texts.pack(side="right", fill="x", expand=True, padx=px(10), pady=px(7))
        tk.Label(texts, text=os.path.basename(path), bg=SURFACE, fg=TEXT,
                 font=(FONT, 9, "bold"), anchor="e").pack(fill="x")
        detail = tk.Label(texts, text=os.path.basename(os.path.dirname(path)), bg=SURFACE, fg=MUTED,
                          font=(FONT, 8), anchor="e", justify="right")
        detail.pack(fill="x")

        remove = tk.Label(row, text="✕", bg=SURFACE, fg=FAINT, font=(FONT, 10), cursor="hand2")
        remove.pack(side="left", padx=(px(10), px(2)))
        remove.bind("<Button-1>", lambda e: self.remove(path))
        remove.bind("<Enter>", lambda e: remove.config(fg="#f87171"))
        remove.bind("<Leave>", lambda e: remove.config(fg=FAINT))
        badge = tk.Label(row, font=(FONT, 8, "bold"), padx=px(8), pady=px(1))
        badge.pack(side="left", padx=px(4))

        texts.bind("<Configure>", lambda e: detail.config(wraplength=max(e.width - 4, px(100))))
        self.rows[path] = {"frame": row, "badge": badge, "detail": detail}
        self._set_status(path, "wait")

    def _set_status(self, path, status, detail=None):
        text, bg, fg = BADGES[status]
        r = self.rows[path]
        r["badge"].config(text=text, bg=bg, fg=fg)
        if detail is not None:
            r["detail"].config(text=detail, fg=fg if status in ("err", "warn") else MUTED)
        self.update_idletasks()

    # --- כפתורים (נבנים לפני הרשימה כדי שיישארו בתחתית)
    def _build_actions(self, parent):
        bar = tk.Frame(parent, bg=BG)
        bar.pack(side="bottom", fill="x", pady=(px(10), 0))
        self.run_btn = RoundButton(bar, "סדר את הקבצים", self.run, width=150, height=38)
        self.run_btn.pack(side="right")
        RoundButton(bar, "פתח תיקייה", self.open_folder, width=104, height=38,
                    style="ghost", font_size=9).pack(side="right", padx=px(8))
        tk.Label(bar, text='הקובץ המסודר נשמר ליד המקור', bg=BG, fg=FAINT,
                 font=(FONT, 8)).pack(side="left")

    # --- פעולות
    def _on_drop(self, event):
        self._set_drop_hover(False)
        self.add_files(self.tk.splitlist(event.data))
        return "copy"

    def pick_files(self):
        start = os.path.dirname(self.files[-1]) if self.files else APP_DIR
        paths = filedialog.askopenfilenames(title="בחר קבצי אקסל לסידור", initialdir=start,
                                            filetypes=[("Excel", "*.xlsx *.xlsm")])
        self.add_files(paths)

    def add_files(self, paths):
        for p in paths:
            p = os.path.normpath(p)
            name = os.path.basename(p)
            if (not p.lower().endswith((".xlsx", ".xlsm")) or name.startswith("~$")
                    or p in self.files):
                continue
            self.files.append(p)
            self._add_row(p)
        self._refresh()

    def remove(self, path):
        self.files.remove(path)
        self.rows.pop(path)["frame"].destroy()
        self._refresh()

    def clear(self):
        for p in list(self.files):
            self.remove(p)

    def _refresh(self):
        self.count_lbl.config(text=str(len(self.files)))
        if self.files:
            self.empty_lbl.pack_forget()
        else:
            self.empty_lbl.pack(pady=px(18))
        self.run_btn.set_enabled(bool(self.files))

    def open_folder(self):
        folder = os.path.dirname(self.files[-1]) if self.files else APP_DIR
        os.startfile(folder)

    def run(self):
        self.run_btn.set_enabled(False)
        ok = 0
        for path in self.files:
            self._set_status(path, "run", os.path.basename(os.path.dirname(path)))
            try:
                status, msg = process_file(path)
            except PermissionError:
                status, msg = "err", "לא ניתן לשמור – סגור את הקובץ המסודר באקסל ונסה שוב"
            except Exception as e:
                traceback.print_exc()
                status, msg = "err", f"שגיאה: {e}"
            if status != "err":
                ok += 1
                msg = f"{msg}\n{os.path.basename(output_path_for(path))}"
            self._set_status(path, status, msg)
        self.run_btn.set_enabled(True)

        failed = len(self.files) - ok
        if failed:
            messagebox.showwarning("הסתיים", f"סודרו {ok} קבצים, ונכשלו {failed}.\nפרטים מופיעים ברשימה.")
        elif len(self.files) == 1:
            if messagebox.askyesno("הסתיים", "הקובץ סודר ונשמר בהצלחה.\nלפתוח אותו עכשיו?"):
                os.startfile(output_path_for(self.files[0]))
        elif messagebox.askyesno("הסתיים", f"כל {ok} הקבצים סודרו בהצלחה.\nלפתוח את התיקייה?"):
            self.open_folder()


if __name__ == "__main__":
    try:
        from ctypes import windll
        windll.shcore.SetProcessDpiAwareness(1)  # טקסט חד במסכים ברזולוציה גבוהה
    except Exception:
        pass
    App().mainloop()
